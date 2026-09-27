from pptx import Presentation
from pptx.util import Inches, Pt, Emu
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE
from pptx.oxml.ns import qn

SLIDE_WIDTH = Inches(13.333)
SLIDE_HEIGHT = Inches(7.5)

DARK_NAVY = RGBColor(0x1B, 0x3A, 0x5C)
ACCENT_ORANGE = RGBColor(0xE8, 0x71, 0x2B)
ACCENT_GREEN = RGBColor(0x2E, 0x7D, 0x32)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
BLACK = RGBColor(0x00, 0x00, 0x00)
GRAY_TEXT = RGBColor(0x33, 0x33, 0x33)
LIGHT_GRAY = RGBColor(0xE0, 0xE0, 0xE0)
BOX_BORDER = RGBColor(0x90, 0xA4, 0xAE)
TABLE_HEADER_BG = RGBColor(0x37, 0x47, 0x4F)

prs = Presentation()
prs.slide_width = SLIDE_WIDTH
prs.slide_height = SLIDE_HEIGHT


def add_bottom_bar(slide):
    bar_h = Inches(0.15)
    bar_y = SLIDE_HEIGHT - bar_h
    half_w = int(SLIDE_WIDTH / 2)

    left_bar = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, bar_y, half_w, bar_h)
    left_bar.fill.solid()
    left_bar.fill.fore_color.rgb = ACCENT_ORANGE
    left_bar.line.fill.background()

    right_bar = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, half_w, bar_y, half_w, bar_h)
    right_bar.fill.solid()
    right_bar.fill.fore_color.rgb = ACCENT_GREEN
    right_bar.line.fill.background()


def add_page_number(slide, num):
    txBox = slide.shapes.add_textbox(SLIDE_WIDTH - Inches(0.8), SLIDE_HEIGHT - Inches(0.5), Inches(0.6), Inches(0.3))
    tf = txBox.text_frame
    p = tf.paragraphs[0]
    p.text = str(num)
    p.font.size = Pt(11)
    p.font.color.rgb = ACCENT_ORANGE
    p.alignment = PP_ALIGN.RIGHT


def set_bold_run(paragraph, text, font_size, color, bold=True):
    run = paragraph.add_run()
    run.text = text
    run.font.size = font_size
    run.font.bold = bold
    run.font.color.rgb = color
    return run


def add_bullet(text_frame, text, level=0, font_size=Pt(14), bold=False, color=GRAY_TEXT, spacing_before=Pt(4)):
    p = text_frame.add_paragraph()
    p.level = level
    p.space_before = spacing_before
    run = p.add_run()
    run.text = text
    run.font.size = font_size
    run.font.bold = bold
    run.font.color.rgb = color
    return p


def add_mixed_bullet(text_frame, bold_part, normal_part, level=0, font_size=Pt(14)):
    p = text_frame.add_paragraph()
    p.level = level
    p.space_before = Pt(4)
    r1 = p.add_run()
    r1.text = bold_part
    r1.font.size = font_size
    r1.font.bold = True
    r1.font.color.rgb = GRAY_TEXT
    r2 = p.add_run()
    r2.text = normal_part
    r2.font.size = font_size
    r2.font.bold = False
    r2.font.color.rgb = GRAY_TEXT
    return p


def add_bordered_box(slide, left, top, width, height):
    box = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, left, top, width, height)
    box.fill.background()
    box.line.color.rgb = BOX_BORDER
    box.line.width = Pt(1.5)
    return box


# ============================================================
# SLIDE 1 - TITLE
# ============================================================
slide1 = prs.slides.add_slide(prs.slide_layouts[6])

title_box = slide1.shapes.add_textbox(Inches(0.5), Inches(0.3), Inches(10), Inches(1.0))
tf = title_box.text_frame
tf.word_wrap = True
p = tf.paragraphs[0]
run = p.add_run()
run.text = "SMART INDIA HACKATHON 2026"
run.font.size = Pt(36)
run.font.bold = True
run.font.color.rgb = ACCENT_ORANGE

content_box = slide1.shapes.add_textbox(Inches(0.8), Inches(1.5), Inches(8), Inches(5.5))
tf = content_box.text_frame
tf.word_wrap = True

p = tf.paragraphs[0]
r1 = p.add_run()
r1.text = "Problem Statement ID "
r1.font.size = Pt(18)
r1.font.bold = True
r1.font.color.rgb = BLACK
r1b = p.add_run()
r1b.text = "- "
r1b.font.size = Pt(18)
r1b.font.color.rgb = BLACK
r1c = p.add_run()
r1c.text = "SIH26035"
r1c.font.size = Pt(18)
r1c.font.bold = True
r1c.font.color.rgb = ACCENT_ORANGE

p2 = tf.add_paragraph()
p2.space_before = Pt(10)
r2a = p2.add_run()
r2a.text = "Problem Statement Title- "
r2a.font.size = Pt(18)
r2a.font.bold = True
r2a.font.color.rgb = BLACK
r2b = p2.add_run()
r2b.text = ("Development of a Software Program/Application for Generation of Test Reports "
            "for Non-Automatic Weighing Instruments (NAWI) as per OIML Recommendation R-76.")
r2b.font.size = Pt(16)
r2b.font.bold = False
r2b.font.color.rgb = ACCENT_ORANGE

p3 = tf.add_paragraph()
p3.space_before = Pt(14)
r3a = p3.add_run()
r3a.text = "Theme- "
r3a.font.size = Pt(18)
r3a.font.bold = True
r3a.font.color.rgb = BLACK
r3b = p3.add_run()
r3b.text = "Smart Automation"
r3b.font.size = Pt(18)
r3b.font.color.rgb = ACCENT_ORANGE

p4 = tf.add_paragraph()
p4.space_before = Pt(8)
r4a = p4.add_run()
r4a.text = "PS Category- "
r4a.font.size = Pt(18)
r4a.font.bold = True
r4a.font.color.rgb = BLACK
r4b = p4.add_run()
r4b.text = "Software"
r4b.font.size = Pt(18)
r4b.font.color.rgb = ACCENT_ORANGE

p5 = tf.add_paragraph()
p5.space_before = Pt(8)
r5a = p5.add_run()
r5a.text = "Organization- "
r5a.font.size = Pt(18)
r5a.font.bold = True
r5a.font.color.rgb = BLACK
r5b = p5.add_run()
r5b.text = "Department of Consumer Affairs, Ministry of Consumer Affairs"
r5b.font.size = Pt(18)
r5b.font.color.rgb = ACCENT_ORANGE


# ============================================================
# SLIDE 2 - PROPOSED SOLUTION
# ============================================================
slide2 = prs.slides.add_slide(prs.slide_layouts[6])

hdr = slide2.shapes.add_textbox(Inches(0.5), Inches(0.3), Inches(9), Inches(0.8))
tf = hdr.text_frame
p = tf.paragraphs[0]
run = p.add_run()
run.text = "PROPOSED SOLUTION"
run.font.size = Pt(32)
run.font.bold = True
run.font.color.rgb = DARK_NAVY

body = slide2.shapes.add_textbox(Inches(0.7), Inches(1.2), Inches(7.5), Inches(5.5))
tf = body.text_frame
tf.word_wrap = True

items = [
    ("Instrument Data Capture: ", "Structured digital forms to record manufacturer details, model number, serial number, accuracy class, and all metrological parameters."),
    ("Environmental Logging: ", "Records temperature, humidity, and atmospheric pressure for each test session as required by OIML R-76."),
    ("Automated MPE Calculation: ", "The system computes Maximum Permissible Errors based on the load, verification scale interval (e), and accuracy class of the instrument."),
    ("OIML R-76 Test Modules: ", "Dedicated tabs for Repeatability, Eccentricity, Weighing Performance, and Tare tests with real-time Pass/Fail compliance checks."),
    ("Report Generation: ", "Produces standardized, print-ready test reports in PDF format with all instrument and laboratory details auto-populated."),
    ("Report Repository: ", "All completed reports are stored digitally with instrument-wise search and retrieval, maintaining a full test history."),
]

for i, (bld, nrm) in enumerate(items):
    if i == 0:
        p = tf.paragraphs[0]
    else:
        p = tf.add_paragraph()
    p.space_before = Pt(6)
    p.level = 0
    num_run = p.add_run()
    num_run.text = f"{i+1}.  "
    num_run.font.size = Pt(14)
    num_run.font.bold = True
    num_run.font.color.rgb = GRAY_TEXT
    b_run = p.add_run()
    b_run.text = bld
    b_run.font.size = Pt(14)
    b_run.font.bold = True
    b_run.font.color.rgb = GRAY_TEXT
    n_run = p.add_run()
    n_run.text = nrm
    n_run.font.size = Pt(13)
    n_run.font.bold = False
    n_run.font.color.rgb = GRAY_TEXT

add_bottom_bar(slide2)
add_page_number(slide2, 2)


# ============================================================
# SLIDE 3 - TECHNICAL APPROACH
# ============================================================
slide3 = prs.slides.add_slide(prs.slide_layouts[6])

hdr = slide3.shapes.add_textbox(Inches(0.5), Inches(0.3), Inches(9), Inches(0.8))
tf = hdr.text_frame
p = tf.paragraphs[0]
run = p.add_run()
run.text = "TECHNICAL APPROACH"
run.font.size = Pt(32)
run.font.bold = True
run.font.color.rgb = DARK_NAVY

# TECH STACK section
ts_hdr = slide3.shapes.add_textbox(Inches(0.7), Inches(1.1), Inches(5), Inches(0.4))
tf = ts_hdr.text_frame
p = tf.paragraphs[0]
r = p.add_run()
r.text = "TECH STACK"
r.font.size = Pt(18)
r.font.bold = True
r.font.color.rgb = DARK_NAVY

ts_body = slide3.shapes.add_textbox(Inches(0.9), Inches(1.5), Inches(5), Inches(2.0))
tf = ts_body.text_frame
tf.word_wrap = True
stack_items = [
    "Next.js 16 (React) with TypeScript",
    "Tailwind CSS",
    "SQLite Database (Prisma ORM)",
    "Next.js API Routes (Serverless)",
    "html2pdf.js / Browser Print API",
    "Vercel or any Node.js host"
]
for i, item in enumerate(stack_items):
    if i == 0:
        p = tf.paragraphs[0]
    else:
        p = tf.add_paragraph()
    p.space_before = Pt(3)
    r = p.add_run()
    r.text = "\u2022  " + item
    r.font.size = Pt(13)
    r.font.color.rgb = GRAY_TEXT

# INPUTS REQUIRED section
ir_hdr = slide3.shapes.add_textbox(Inches(0.7), Inches(3.7), Inches(5), Inches(0.4))
tf = ir_hdr.text_frame
p = tf.paragraphs[0]
r = p.add_run()
r.text = "INPUTS REQUIRED"
r.font.size = Pt(18)
r.font.bold = True
r.font.color.rgb = DARK_NAVY

ir_body = slide3.shapes.add_textbox(Inches(0.9), Inches(4.1), Inches(5.5), Inches(2.5))
tf = ir_body.text_frame
tf.word_wrap = True
inputs = [
    ("Instrument specification sheet: ", "capacity, class, scale intervals"),
    ("Laboratory conditions: ", "temperature, humidity, barometric pressure"),
    ("Calibrated test weights ", "for each load point"),
    ("OIML R-76 reference tables ", "for MPE thresholds"),
]
for i, (bld, nrm) in enumerate(inputs):
    if i == 0:
        p = tf.paragraphs[0]
    else:
        p = tf.add_paragraph()
    p.space_before = Pt(3)
    r0 = p.add_run()
    r0.text = "\u2022  "
    r0.font.size = Pt(13)
    r0.font.color.rgb = GRAY_TEXT
    rb = p.add_run()
    rb.text = bld
    rb.font.size = Pt(13)
    rb.font.bold = True
    rb.font.color.rgb = GRAY_TEXT
    rn = p.add_run()
    rn.text = nrm
    rn.font.size = Pt(13)
    rn.font.color.rgb = GRAY_TEXT

# Right side: COMPARISON TABLE
comp_hdr = slide3.shapes.add_textbox(Inches(6.8), Inches(1.1), Inches(6), Inches(0.4))
tf = comp_hdr.text_frame
p = tf.paragraphs[0]
r = p.add_run()
r.text = "MANUAL vs. AUTOMATED REPORT GENERATION"
r.font.size = Pt(14)
r.font.bold = True
r.font.color.rgb = DARK_NAVY

rows, cols = 6, 3
tbl = slide3.shapes.add_table(rows, cols, Inches(6.8), Inches(1.6), Inches(6.0), Inches(3.0)).table
tbl.columns[0].width = Inches(1.5)
tbl.columns[1].width = Inches(2.2)
tbl.columns[2].width = Inches(2.3)

headers = ["Aspect", "Manual Process", "Our Application"]
data = [
    ["Data Entry", "Spreadsheets, copy-paste", "Structured digital forms"],
    ["Calculations", "Manual formula application", "Auto MPE computation"],
    ["Compliance", "Visual comparison by tester", "Real-time Pass/Fail"],
    ["Report Format", "Inconsistent templates", "Standardized PDF output"],
    ["History", "Scattered files on disk", "Searchable database"],
]

for ci, h in enumerate(headers):
    cell = tbl.cell(0, ci)
    cell.text = h
    for paragraph in cell.text_frame.paragraphs:
        paragraph.font.size = Pt(11)
        paragraph.font.bold = True
        paragraph.font.color.rgb = WHITE
    cell.fill.solid()
    cell.fill.fore_color.rgb = TABLE_HEADER_BG

for ri, row_data in enumerate(data):
    for ci, val in enumerate(row_data):
        cell = tbl.cell(ri + 1, ci)
        cell.text = val
        for paragraph in cell.text_frame.paragraphs:
            paragraph.font.size = Pt(11)
            paragraph.font.color.rgb = GRAY_TEXT

add_bottom_bar(slide3)
add_page_number(slide3, 3)


# ============================================================
# SLIDE 4 - MODEL APPROACH
# ============================================================
slide4 = prs.slides.add_slide(prs.slide_layouts[6])

hdr = slide4.shapes.add_textbox(Inches(0.5), Inches(0.3), Inches(9), Inches(0.8))
tf = hdr.text_frame
p = tf.paragraphs[0]
run = p.add_run()
run.text = "MODEL APPROACH"
run.font.size = Pt(32)
run.font.bold = True
run.font.color.rgb = DARK_NAVY

# 5 pipeline boxes
box_w = Inches(2.3)
box_h = Inches(2.0)
box_y = Inches(1.3)
gap = Inches(0.2)
start_x = Inches(0.5)

steps = [
    ("1", "Register Instrument", "Enter manufacturer, model, serial number, accuracy class, scale intervals (e, d), Min/Max capacity."),
    ("2", "Configure Test Session", "Select instrument, assign tester, log ambient temperature, humidity, and barometric pressure."),
    ("3", "Record Observations", "Input load and indication readings for Repeatability, Eccentricity, Weighing, and Tare tests."),
    ("4", "Auto-Calculate & Validate", "System computes error, determines MPE from OIML R-76 tables, and flags each reading Pass or Fail."),
    ("5", "Generate Report", "Mark report complete; system aggregates results, determines overall compliance, and outputs a standardized PDF."),
]

for i, (num, title, desc) in enumerate(steps):
    x = start_x + i * (box_w + gap)
    add_bordered_box(slide4, x, box_y, box_w, box_h)

    num_box = slide4.shapes.add_textbox(x + Inches(0.15), box_y + Inches(0.1), Inches(0.4), Inches(0.35))
    tf = num_box.text_frame
    p = tf.paragraphs[0]
    r = p.add_run()
    r.text = num
    r.font.size = Pt(20)
    r.font.bold = True
    r.font.color.rgb = DARK_NAVY

    title_box = slide4.shapes.add_textbox(x + Inches(0.15), box_y + Inches(0.4), box_w - Inches(0.3), Inches(0.45))
    tf = title_box.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    r = p.add_run()
    r.text = title
    r.font.size = Pt(13)
    r.font.bold = True
    r.font.color.rgb = GRAY_TEXT

    desc_box = slide4.shapes.add_textbox(x + Inches(0.15), box_y + Inches(0.85), box_w - Inches(0.3), box_h - Inches(1.0))
    tf = desc_box.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    r = p.add_run()
    r.text = desc
    r.font.size = Pt(11)
    r.font.color.rgb = GRAY_TEXT

# Bottom left: CORE CALCULATION LOGIC
cl_hdr = slide4.shapes.add_textbox(Inches(0.7), Inches(3.6), Inches(5), Inches(0.4))
tf = cl_hdr.text_frame
p = tf.paragraphs[0]
r = p.add_run()
r.text = "CORE CALCULATION LOGIC"
r.font.size = Pt(18)
r.font.bold = True
r.font.color.rgb = DARK_NAVY

cl_body = slide4.shapes.add_textbox(Inches(0.9), Inches(4.05), Inches(5), Inches(3.0))
tf = cl_body.text_frame
tf.word_wrap = True
calcs = [
    "Error = Indication - Applied Load",
    "MPE determined by load range and accuracy class",
    "Class III: 0.5e (0-500e), 1.0e (500-2000e), 1.5e (>2000e)",
    "Class II: 0.5e (0-5000e), 1.0e (5000-20000e), 1.5e (>20000e)",
    "Result = Pass if |Error| <= MPE, else Fail",
]
for i, item in enumerate(calcs):
    if i == 0:
        p = tf.paragraphs[0]
    else:
        p = tf.add_paragraph()
    p.space_before = Pt(3)
    r = p.add_run()
    r.text = "\u2022  " + item
    r.font.size = Pt(13)
    r.font.color.rgb = GRAY_TEXT

# Bottom right: KEY SOFTWARE MODULES
km_hdr = slide4.shapes.add_textbox(Inches(6.8), Inches(3.6), Inches(6), Inches(0.4))
tf = km_hdr.text_frame
p = tf.paragraphs[0]
r = p.add_run()
r.text = "KEY SOFTWARE MODULES"
r.font.size = Pt(18)
r.font.bold = True
r.font.color.rgb = DARK_NAVY

km_body = slide4.shapes.add_textbox(Inches(7.0), Inches(4.05), Inches(5.5), Inches(3.0))
tf = km_body.text_frame
tf.word_wrap = True
modules = [
    "Instrument Registry (CRUD with specifications)",
    "Test Session Manager (environmental conditions)",
    "Observation Entry per OIML R-76 test type",
    "MPE Calculator Engine (oimlHelpers.ts)",
    "PDF Report Builder with auto-populated fields",
    "Dashboard for monitoring and search",
]
for i, item in enumerate(modules):
    if i == 0:
        p = tf.paragraphs[0]
    else:
        p = tf.add_paragraph()
    p.space_before = Pt(3)
    r = p.add_run()
    r.text = "\u2022  " + item
    r.font.size = Pt(13)
    r.font.color.rgb = GRAY_TEXT

add_bottom_bar(slide4)
add_page_number(slide4, 4)


# ============================================================
# SLIDE 5 - FEASIBILITY, VIABILITY, IMPACT AND BENEFITS
# ============================================================
slide5 = prs.slides.add_slide(prs.slide_layouts[6])

hdr = slide5.shapes.add_textbox(Inches(0.5), Inches(0.3), Inches(11), Inches(0.8))
tf = hdr.text_frame
p = tf.paragraphs[0]
run = p.add_run()
run.text = "FEASIBILITY, VIABILITY, IMPACT AND BENEFITS"
run.font.size = Pt(30)
run.font.bold = True
run.font.color.rgb = DARK_NAVY

quad_w = Inches(5.8)
quad_h = Inches(2.5)
left_x = Inches(0.6)
right_x = Inches(6.8)
top_y = Inches(1.3)
bot_y = Inches(4.1)

quads = [
    (left_x, top_y, "Feasibility", [
        "Built with well-established open-source tools (Next.js, Prisma, SQLite) that require no licensing cost.",
        "SQLite runs embedded, so no separate database server is needed for deployment.",
        "OIML R-76 MPE tables are well-defined and straightforward to implement in code.",
        "Can be deployed on any machine with Node.js installed, including air-gapped lab networks.",
    ]),
    (right_x, top_y, "Viability", [
        "Zero infrastructure cost for small labs; scales to cloud hosting if needed.",
        "Reduces report generation time from hours to minutes.",
        "Eliminates duplicate data entry across multiple spreadsheet templates.",
        "Architecture supports future OIML recommendation revisions without rewriting core logic.",
    ]),
    (left_x, bot_y, "Impact", [
        "Ensures uniformity and consistency across all designated testing laboratories.",
        "Removes manual calculation errors that can delay model approvals.",
        "Creates a searchable digital archive of all test reports for auditing.",
        "Aligns India's NAWI testing workflow with international best practices.",
    ]),
    (right_x, bot_y, "Benefits", [
        "Lab technicians spend less time on paperwork, more on actual testing.",
        "Standardized output format simplifies review by approval authorities.",
        "Instrument-wise test history enables trend analysis and quality tracking.",
        "Dashboard gives lab managers real-time visibility into report pipeline status.",
    ]),
]

for (x, y, title, bullets) in quads:
    add_bordered_box(slide5, x, y, quad_w, quad_h)

    t_box = slide5.shapes.add_textbox(x + Inches(0.2), y + Inches(0.1), quad_w - Inches(0.4), Inches(0.35))
    tf = t_box.text_frame
    p = tf.paragraphs[0]
    r = p.add_run()
    r.text = title
    r.font.size = Pt(16)
    r.font.bold = True
    r.font.color.rgb = DARK_NAVY

    b_box = slide5.shapes.add_textbox(x + Inches(0.3), y + Inches(0.5), quad_w - Inches(0.5), quad_h - Inches(0.6))
    tf = b_box.text_frame
    tf.word_wrap = True
    for i, bullet in enumerate(bullets):
        if i == 0:
            p = tf.paragraphs[0]
        else:
            p = tf.add_paragraph()
        p.space_before = Pt(3)
        r = p.add_run()
        r.text = "\u2022  " + bullet
        r.font.size = Pt(12)
        r.font.color.rgb = GRAY_TEXT

add_bottom_bar(slide5)
add_page_number(slide5, 5)


# ============================================================
# SLIDE 6 - REFERENCES & RESEARCH
# ============================================================
slide6 = prs.slides.add_slide(prs.slide_layouts[6])

hdr = slide6.shapes.add_textbox(Inches(0.5), Inches(0.3), Inches(9), Inches(0.8))
tf = hdr.text_frame
p = tf.paragraphs[0]
run = p.add_run()
run.text = "REFERENCES AND RESEARCH"
run.font.size = Pt(32)
run.font.bold = True
run.font.color.rgb = DARK_NAVY

# Primary References
pr_hdr = slide6.shapes.add_textbox(Inches(0.7), Inches(1.2), Inches(11), Inches(0.4))
tf = pr_hdr.text_frame
p = tf.paragraphs[0]
r = p.add_run()
r.text = "PRIMARY REFERENCES"
r.font.size = Pt(18)
r.font.bold = True
r.font.color.rgb = DARK_NAVY

pr_body = slide6.shapes.add_textbox(Inches(0.9), Inches(1.65), Inches(11), Inches(2.2))
tf = pr_body.text_frame
tf.word_wrap = True

primary_refs = [
    'OIML R 76-1 (Edition 2006) "Non-automatic weighing instruments, Part 1: Metrological and technical requirements - Tests"',
    'OIML R 76-2 (Edition 2007) "Non-automatic weighing instruments, Part 2: Test report format"',
    "The Legal Metrology Act, 2009 (Act No. 1 of 2010), Government of India",
    "The Legal Metrology (General) Rules, 2011, Ministry of Consumer Affairs, Food and Public Distribution",
    "IS 9281:2023 (Bureau of Indian Standards) - specification for Electronic Weighing Systems",
]
for i, ref in enumerate(primary_refs):
    if i == 0:
        p = tf.paragraphs[0]
    else:
        p = tf.add_paragraph()
    p.space_before = Pt(4)
    r = p.add_run()
    r.text = f"{i+1}.  {ref}"
    r.font.size = Pt(13)
    r.font.color.rgb = GRAY_TEXT

# Technical Resources
tr_hdr = slide6.shapes.add_textbox(Inches(0.7), Inches(4.0), Inches(11), Inches(0.4))
tf = tr_hdr.text_frame
p = tf.paragraphs[0]
r = p.add_run()
r.text = "TECHNICAL RESOURCES"
r.font.size = Pt(18)
r.font.bold = True
r.font.color.rgb = DARK_NAVY

tr_body = slide6.shapes.add_textbox(Inches(0.9), Inches(4.45), Inches(11), Inches(2.0))
tf = tr_body.text_frame
tf.word_wrap = True

tech_refs = [
    "consumeraffairs.gov.in - Department of Consumer Affairs, Legal Metrology Division",
    "oiml.org - International Organization of Legal Metrology, publications and recommendations",
    "Next.js Documentation (nextjs.org/docs) - App Router, API Routes, Server Components",
    "Prisma ORM Documentation (prisma.io/docs) - SQLite integration, schema design",
    "WELMEC Guide 2 (2019) - Directive 2014/31/EU on Non-automatic Weighing Instruments (EU context for OIML R 76)",
]
for i, ref in enumerate(tech_refs):
    if i == 0:
        p = tf.paragraphs[0]
    else:
        p = tf.add_paragraph()
    p.space_before = Pt(4)
    r = p.add_run()
    r.text = f"{i+1}.  {ref}"
    r.font.size = Pt(13)
    r.font.color.rgb = GRAY_TEXT

add_bottom_bar(slide6)
add_page_number(slide6, 6)


prs.save("SIH26035_Presentation.pptx")
print("Presentation generated successfully.")
